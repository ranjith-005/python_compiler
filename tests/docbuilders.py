"""Build real PDF and PPTX bytes for the module-upload tests.

The uploader is only worth testing against the formats it actually receives, so
these produce genuine files: the deck goes through python-pptx, and the PDF is a
small hand-written document with an uncompressed text stream that pypdf reads
back the same way it reads any other PDF.
"""

from __future__ import annotations

import io


def deck_bytes(slides: list[tuple[str, list[str]]]) -> bytes:
    """A .pptx with one slide per (title, lines) pair."""
    from pptx import Presentation

    presentation = Presentation()
    layout = presentation.slide_layouts[1]          # Title and Content
    for title, lines in slides:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        frame = slide.placeholders[1].text_frame
        if lines:
            frame.text = lines[0]
            for line in lines[1:]:
                frame.add_paragraph().text = line
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def pdf_bytes(pages: list[list[str]]) -> bytes:
    """A .pdf with one page per list of lines, in Helvetica."""
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)                          # 1-based object number

    # Reserve 1 (catalog) and 2 (page tree); their bodies need the page ids.
    objects.append(b"")
    objects.append(b"")
    font_id = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    page_ids: list[int] = []
    for lines in pages:
        text = "BT\n/F1 12 Tf\n72 720 Td\n15 TL\n"
        for line in lines:
            text += f"({_escape(line)}) Tj\nT*\n"
        text += "ET"
        stream = text.encode("latin-1", "replace")
        content_id = add(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
        page_id = add(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 " + str(font_id).encode() + b" 0 R >> >> "
            b"/Contents " + str(content_id).encode() + b" 0 R >>"
        )
        page_ids.append(page_id)

    kids = b" ".join(str(pid).encode() + b" 0 R" for pid in page_ids)
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[1] = (
        b"<< /Type /Pages /Kids [" + kids + b"] /Count " + str(len(page_ids)).encode() + b" >>"
    )

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += str(number).encode() + b" 0 obj\n" + body + b"\nendobj\n"

    xref_at = len(out)
    out += b"xref\n0 " + str(len(objects) + 1).encode() + b"\n"
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        b"trailer\n<< /Size " + str(len(objects) + 1).encode() + b" /Root 1 0 R >>\n"
        b"startxref\n" + str(xref_at).encode() + b"\n%%EOF\n"
    )
    return bytes(out)
